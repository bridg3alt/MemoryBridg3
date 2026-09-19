import math, re, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
from pptx import Presentation
from pptx.util import Emu

p = Presentation("MemoryBridge_IEDC_Summit_2026.pptx")
PLACE = re.compile(r"\[|lorem|ipsum|TODO|Add details|Member \d|Feature 0|What does it do", re.I)
issues = []
for i, s in enumerate(p.slides, 1):
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        tf = sh.text_frame
        txt = tf.text
        if not txt.strip():
            continue
        if PLACE.search(txt):
            issues.append(f"slide {i:02d} {sh.name:9s} PLACEHOLDER: {txt[:70]!r}")
        # effective font size
        sizes = [r.font.size.pt for para in tf.paragraphs for r in para.runs if r.font.size]
        if not sizes:
            xml = sh._element.xml
            m = re.findall(r'sz="(\d+)"', xml)
            sizes = [int(m[0]) / 100] if m else [18]
        size = max(sizes)
        w_pt = Emu(sh.width).inches * 72
        h_pt = Emu(sh.height).inches * 72
        cpl = max(1, int(w_pt / (size * 0.5)))
        lines = sum(max(1, math.ceil(len(line) / cpl)) for line in txt.split("\n"))
        need = lines * size * 1.22
        if need > h_pt * 1.35:
            issues.append(f"slide {i:02d} {sh.name:9s} OVERFLOW?: {size:.0f}pt box {Emu(sh.width).inches:.1f}x{Emu(sh.height).inches:.1f}in "
                          f"needs ~{lines} lines ({need:.0f}pt vs {h_pt:.0f}pt) :: {txt[:60]!r}")
print("\n".join(issues) if issues else "no issues found")
