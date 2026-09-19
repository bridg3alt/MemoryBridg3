"""Fill the INCLUCODE template with MemoryBridge content.

Run:  python presentation/fill_deck.py
In:   presentation/work2.pptx   (template + 2 duplicated slides)
Out:  presentation/MemoryBridge_IEDC_Summit_2026.pptx
"""
import copy
import re
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "work2.pptx"
TARGET = HERE / "MemoryBridge_IEDC_Summit_2026.pptx"

NAME = "[Your Name]"
COLLEGE = "[Your College]"


def set_text(shape, text):
    """Replace a shape's text, one paragraph per line, keeping the template's formatting."""
    frame = shape.text_frame
    lines = text.split("\n")
    paragraphs = frame.paragraphs

    while len(paragraphs) < len(lines):
        clone = copy.deepcopy(paragraphs[-1]._p)
        paragraphs[-1]._p.addnext(clone)
        paragraphs = frame.paragraphs
    for extra in list(paragraphs)[len(lines):]:
        extra._p.getparent().remove(extra._p)
    paragraphs = frame.paragraphs

    for paragraph, line in zip(paragraphs, lines):
        runs = paragraph.runs
        if not runs:
            paragraph.add_run()
            runs = paragraph.runs
        runs[0].text = line
        for spare in runs[1:]:
            spare._r.getparent().remove(spare._r)


# slide number (1-based, after the two duplications) -> {shape name: new text}
CONTENT = {
    # ---------------------------------------------------------------- 1 title
    1: {
        "Text 4": "DATA SCIENCE CONTEST",
        "Text 5": "IEDC SUMMIT 2026",
        "Text 6": "Music as a cue for personal memory",
        "Text 11": f"MemoryBridge\n{NAME}\n{COLLEGE}",
    },
    # ------------------------------------------------------------- 2 overview
    2: {
        "Text 5": "Project Overview",
        "Text 6": "One submission, four workstreams: data preparation, modelling, evaluation and the app.",
        "Text 14": "SUBMITTED BY",
        "Text 17": NAME,
        "Text 18": "PROJECT TITLE",
        "Text 21": "MemoryBridge — music as a cue for personal memory",
        "Text 22": "WORK BREAKDOWN",
        "Text 25": "TurboPrep\nCleaning, joining and standardising 9,840 songs",
        "Text 27": "RapidMiner\nSeven features, three models, cross validation, batch scoring",
        "Text 29": "Python\nMusicBrainz API, simulation, Streamlit app, SQLite",
        "Text 31": "Evaluation\nHoldout AUC 0.804, top-10 hit rate 74.2%",
    },
    # -------------------------------------------------------------- 3 problem
    3: {
        "Text 6": "Older listeners cannot find the music of their own youth, and the memories it carries are never written down.",
        "Text 14": "WHAT IS THE PROBLEM?",
        "Text 17": (
            "Music heard between the ages of 10 and 30 is the strongest cue for personal memory, "
            "but nobody holds that list of songs for a specific person. Existing playlist tools need "
            "a caregiver who already knows the songs, and when a memory does arrive it is lost again "
            "because nothing captures it."
        ),
        "Text 18": "WHO FACES IT?",
        "Text 21": (
            "Older adults and the families who want their stories\n"
            "Malayalam and Hindi listeners in Kerala and India\n"
            "Barriers: English-only catalogues, no personal song list, no record of the memory"
        ),
        "Text 22": "EVIDENCE / OBSERVATION",
        "Text 25": (
            "Cuddy, Sikka et al. (2021): recollection of popular music peaks in the teens and twenties — the reminiscence bump.\n"
            "In the Million Song Dataset only 1,210 songs with a year and a genre come from the 1950s, against 128,641 from the 2000s."
        ),
    },
    # --------------------------------------------------------------- 4 users
    4: {
        "Text 6": "Who this is built for, what they do today, and the limits the design has to respect.",
        "Text 15": (
            "User Persona\n"
            "Thomas, 66, born 1960, Kerala\n"
            "Listens to Malayalam film songs and some English music\n"
            "Types slowly, speaks easily, uses a phone with help\n"
            "Wants his children to know the stories behind his songs"
        ),
        "Text 17": (
            "Current Journey\n"
            "Tries to recall a song he half remembers\n"
            "Searches in English and gets recent hits\n"
            "Gives up, or finds it and tells nobody\n"
            "The memory arrives and then disappears"
        ),
        "Text 19": (
            "Design Constraints\n"
            "Malayalam and Hindi titles must survive encoding\n"
            "No streaming account, low bandwidth\n"
            "Speech instead of typing\n"
            "No medical claim, no invented memories"
        ),
    },
    # ------------------------------------------------------------ 5 solution
    5: {
        "Text 6": "What the product does, in one line anyone can understand.",
        "Text 15": (
            "ONE-LINE SOLUTION\n"
            "MemoryBridge helps an older listener revisit their own life by suggesting the songs of their youth "
            "and keeping the memory each song brings back."
        ),
        "Text 16": "CORE VALUE",
        "Text 19": (
            "Songs from the person's own language and era, not a generic decade playlist, "
            "and the story they tell is saved and can be reopened by life period."
        ),
        "Text 20": "WHAT YOU BUILT",
        "Text 23": (
            "A Streamlit app over a RapidMiner Random Forest, an Altair AI Studio TurboPrep pipeline "
            "and a local SQLite store."
        ),
        "Text 24": "EXPECTED OUTCOME",
        "Text 27": (
            "Top-10 hit rate 74.2% against 63.2% for the formula alone, "
            "100% language match in the top 10, memories kept by life period."
        ),
    },
    # -------------------------------------------------------- 6 architecture
    6: {
        "Text 6": "How the solution works internally, from raw data to a saved memory.",
        "Text 15": "Input\nBirth year, language, genres",
        "Text 18": "TurboPrep\nClean, join, standardise",
        "Text 21": "RapidMiner\nRandom Forest, 7 features",
        "Text 24": "Output\nTop 10 songs, ranked",
        "Text 27": "Impact\nMemory saved by life period",
        "Text 29": (
            "Five public sources → Python sampling and the MusicBrainz API → TurboPrep → songs_master.csv, 9,840 songs.\n"
            "Python builds 9,589 simulated training rows → RapidMiner cross validation split on user_id → Random Forest → 265,680 predictions.\n"
            "The app reads that file: the baseline picks 200 candidates, the model ranks them, SQLite keeps the memory."
        ),
    },
    # ------------------------------------------------- 7 data (duplicated 9)
    7: {
        "Text 5": "Data & Preparation",
        "Text 6": "Five public datasets, cleaned and joined into one master table of 9,840 songs.",
        "Text 15": "Sources",
        "Text 16": (
            "Million Song Dataset — 1,000,000 tracks\n"
            "tagtraum genres — 280,831\n"
            "Hindi film metadata — 57,005\n"
            "Malayalam songs — 26,232\n"
            "MusicBrainz artist API"
        ),
        "Text 17": "energy is 0 for all 1,000,000 songs, so it was measured and dropped.",
        "Text 19": "Python Preparation",
        "Text 20": (
            "Sample of 928 songs per decade, 1950–2010, seed 42\n"
            "MusicBrainz enrichment at 1 request/second\n"
            "3,502 of 3,525 artists found, language 99.8%"
        ),
        "Text 21": "Stratified because the 1950s hold 1,210 songs against 128,641 in the 2000s.",
        "Text 23": "TurboPrep (AI Studio)",
        "Text 24": (
            "A · Hindi cleaned — 2,000 rows\n"
            "B · Malayalam cleaned — 1,349 rows\n"
            "C · MSD + MusicBrainz join — 6,491 rows\n"
            "D · Append and standardise 15 language codes"
        ),
        "Text 25": "'#' silently broke 5 rows and the system encoding turned Malayalam into '???'. Both found and fixed.",
        "Text 27": "Master Table",
        "Text 28": (
            "songs_master.csv — 9,840 × 14\n"
            "Years 1950–2010, 0 duplicate ids\n"
            "English 5,921 · Hindi 2,000 · Malayalam 1,349\n"
            "Popularity missing in 5,238 rows"
        ),
        "Text 29": "Missing popularity is filled with the training average, 63.633, everywhere it is used.",
    },
    # --------------------------------------------- 8 models (duplicated 11)
    8: {
        "Text 5": "Model Comparison & Validation",
        "Text 6": "Three models, tested so that no listener ever appears in both training and testing.",
        "Text 14": "CROSS VALIDATION",
        "Text 17": (
            "Logistic Regression — AUC 0.790\n"
            "Random Forest — AUC 0.831\n"
            "Gradient Boosted Trees — AUC 0.835\n"
            "Always predicting no — 57.0%"
        ),
        "Text 18": "WHY RANDOM FOREST",
        "Text 21": (
            "AUC ties with GBT inside ±0.073\n"
            "Recall 80.5% against 58.3%\n"
            "Missing a song that would have worked is the expensive error\n"
            "Simpler to explain"
        ),
        "Text 22": "HOLDOUT — 60 UNSEEN LISTENERS",
        "Text 25": (
            "AUC 0.804\n"
            "Top-10 hit rate\n"
            "Random Forest 74.2%\n"
            "Baseline formula 63.2%\n"
            "Random order 40.2%"
        ),
        "Text 27": (
            "Cross Validation with split on batch attribute (user_id): every listener is scored by a model that never saw them. "
            "Trained on SIMULATED interactions — the pipeline is proven, the idea is not yet."
        ),
    },
    # -------------------------------------------------------- 9 key features
    9: {
        "Text 6": "The five things that make this work for the person in front of it.",
        "Text 17": "Two-stage recommender",
        "Text 18": "The baseline picks 200 candidates, the Random Forest ranks them.\n100% language match in the top 10.",
        "Text 22": "Songs in the listener's own language",
        "Text 23": "Malayalam and Hindi are first-class, not an afterthought.\n3,349 rows carry no audio fields and still rank.",
        "Text 27": "Speak the memory",
        "Text 28": "Speech to text in the listener's own language.\nThe transcript is editable; nothing saves until they confirm.",
        "Text 32": "The app says who ranked",
        "Text 33": "Every list states RapidMiner Random Forest or Baseline formula,\nso the person can see how the songs were chosen.",
    },
    # ----------------------------------------------------- 10 accessibility
    10: {
        "Text 6": "Inclusion is a design decision here, not a feature added at the end.",
        "Text 16": "Accessible by Design",
        "Text 17": "Speech input in Malayalam, Hindi or English; large plain type, no emoji; runs on one laptop with no streaming account.",
        "Text 20": "User Safety & Ethics",
        "Text 21": "Not a medical product and no diagnosis; memories stay in a local SQLite file; the app never invents a memory or rewrites the person's words.",
        "Text 24": "Local Relevance",
        "Text 25": "3,349 Malayalam and Hindi songs; UTF-8 fixed so titles render in their own script; a Devanagari font for Hindi text.",
        "Text 27": "Proof: 9 automated tests pass, and a headless end-to-end run saved a 1976 Malayalam song, its memory and its life period.",
    },
    # -------------------------------------------------------- 11 tech stack
    11: {
        "Text 6": "The tools behind each part of the workflow, and why each one was chosen.",
        "Text 15": "Frontend / UI",
        "Text 16": "Streamlit, five pages, paper-and-ink theme, speech input through streamlit-mic-recorder",
        "Text 17": "Python UI so one person could ship something usable.",
        "Text 19": "Backend / Data",
        "Text 20": "Python 3.13, pandas, SQLite with users, memories and interactions, MusicBrainz API at 1 request/second",
        "Text 21": "SQLite keeps everything local: no account, no server.",
        "Text 23": "AI / ML / Models",
        "Text 24": "Altair AI Studio 2026.1.1: TurboPrep, Random Forest, Logistic Regression, Gradient Boosted Trees, Cross Validation",
        "Text 25": "Random Forest chosen on recall, not on AUC alone.",
        "Text 27": "Scoring",
        "Text 28": "Batch scoring writes predictions_grid.csv, 265,680 rows, which the app loads at startup",
        "Text 29": "Batch scoring is how AI Studio is used in practice.",
    },
    # ------------------------------------------------------ 12 demo evidence
    12: {
        "Text 6": "Proof that the workflow runs and the prototype works.",
        "Text 15": "DRIVE FOLDER\n[Paste the Drive link here]",
        "Text 17": "DEMO VIDEO\n[Paste the YouTube link here]",
        "Text 19": "PROCESS FILES\nBoth .rmp files, in the Drive folder",
        "Text 20": "SCREENSHOTS / PRODUCT IMAGES",
        "Text 23": "Image 1\nThe RapidMiner process: cross validation, holdout, demo and grid scoring",
        "Text 25": "Image 2\nThe app's song list, with the line naming the Random Forest as the ranker",
        "Text 27": "Image 3\nThe timeline: a saved memory filed under its life period",
    },
    # ----------------------------------------------------------- 13 impact
    13: {
        "Text 6": "What the numbers show, how they were checked, and where this goes next.",
        "Text 14": "IMPACT METRICS",
        "Text 17": (
            "Top-10 hit rate 74.2% against 63.2%\n"
            "265,680 predictions served to the app\n"
            "9,840 songs across three languages\n"
            "27 profile cells covering birth years 1935–1975"
        ),
        "Text 18": "VALIDATION",
        "Text 21": (
            "Leave-one-listener-out cross validation\n"
            "Holdout of 60 listeners never seen in training\n"
            "Independent scikit-learn check agrees with AI Studio\n"
            "SIMULATED labels — real listeners are the next step"
        ),
        "Text 22": "SCALABILITY",
        "Text 25": (
            "The interactions table already stores real reactions in the training format, so one retrain replaces the simulation.\n"
            "The grid grows by language and birth year, not by user, so more listeners cost nothing."
        ),
        "Text 27": "In one line: MemoryBridge can help an older listener recover the songs of their youth, and keep the memories those songs return.",
    },
    # ---------------------------------------------------------- 14 roadmap
    14: {
        "Text 6": "A practical next-step plan, with the honest gap named first.",
        "Text 15": "Now",
        "Text 17": "• Collect 30–60 real reactions\n• From 3–5 real listeners\n• Replaces the simulated labels",
        "Text 19": "Next 2 Weeks",
        "Text 21": "• Retrain on the real reactions\n• Re-run the holdout test\n• Honest, non-simulated scores",
        "Text 23": "Next 2 Months",
        "Text 25": "• Pilot with families and a care home\n• Memories saved with photos\n• Add Tamil and Telugu to the grid",
        "Text 27": "Support Needed",
        "Text 29": "• Introductions to older listeners\n• Consent and ethics guidance\n• Hosting so scoring can run on a schedule",
        "Text 31": "Contact: [Your Name]  |  [Educational email]  |  [Phone]  |  [YouTube / Drive link]",
    },
    # -------------------------------------------------------------- 15 close
    15: {
        "Text 5": f"{NAME}  •  MemoryBridge",
        "Text 7": "Demo video  ·  Drive folder  ·  docs/project_summary.md",
    },
}


# cards whose first line is a heading rather than body text
BOLD_FIRST_LINE = {
    2: ["Text 25", "Text 27", "Text 29", "Text 31"],
    4: ["Text 15", "Text 17", "Text 19"],
    6: ["Text 15", "Text 18", "Text 21", "Text 24", "Text 27"],
    12: ["Text 23", "Text 25", "Text 27"],
}


def add_fifth_feature(slide):
    """Slide 9 ships with four feature cards. Clone card 3 into a wide fifth row underneath.

    The template has no room for a third 2-column row (the footer starts at 6.83"), so the
    fifth card spans the full width and puts its text beside the title instead of under it.
    """
    by_name = {shape.name: shape for shape in slide.shapes}
    tree = slide.shapes._spTree
    next_id = max(shape.shape_id for shape in slide.shapes) + 1

    # source shape -> (left, top, width, height) in inches for the copy
    layout = {
        "Shape 24": (0.85, 6.00, 11.35, 0.65),   # the card
        "Shape 25": (1.05, 6.06, 0.52, 0.52),    # the coloured circle
        "Text 26": (1.19, 6.18, 0.25, 0.15),     # the number inside it
        "Text 27": (1.75, 6.09, 4.00, 0.25),     # the feature title
        "Text 28": (5.95, 6.11, 6.25, 0.25),     # the description, beside the title
    }
    text = {"Text 26": "5",
            "Text 27": "Memories return by life period",
            "Text 28": "Childhood, College, Early adulthood, Family/Career — each memory is filed under the era it belongs to."}

    for source_name, (left, top, width, height) in layout.items():
        element = copy.deepcopy(by_name[source_name]._element)
        tree.append(element)
        copied = slide.shapes[-1]
        copied._element.nvSpPr.cNvPr.set("id", str(next_id))
        copied._element.nvSpPr.cNvPr.set("name", f"{source_name} copy")
        next_id += 1
        copied.left, copied.top = Inches(left), Inches(top)
        copied.width, copied.height = Inches(width), Inches(height)
        if source_name in text:
            set_text(copied, text[source_name])
            if source_name == "Text 28":  # the body is bold in card 3's first line only
                for run in copied.text_frame.paragraphs[0].runs:
                    run.font.bold = False


def main():
    presentation = Presentation(SOURCE)
    slides = list(presentation.slides)
    if len(slides) != 15:
        raise SystemExit(f"expected 15 slides, found {len(slides)}")

    missing = []
    for index, slide in enumerate(slides, 1):
        by_name = {shape.name: shape for shape in slide.shapes if shape.has_text_frame}
        for name, text in CONTENT.get(index, {}).items():
            if name not in by_name:
                missing.append(f"slide {index}: no shape named {name!r}")
                continue
            set_text(by_name[name], text)

        for name in BOLD_FIRST_LINE.get(index, []):
            first = by_name[name].text_frame.paragraphs[0]
            for run in first.runs:
                run.font.bold = True

        if index == 9:
            add_fifth_feature(slide)

        # the page number in the bottom-right corner
        for shape in slide.shapes:
            if shape.has_text_frame and re.fullmatch(r"\d{2}", shape.text_frame.text.strip() or "x"):
                if shape.left > 11_000_000:
                    set_text(shape, f"{index:02d}")

    if missing:
        raise SystemExit("\n".join(missing))

    presentation.save(TARGET)
    print(f"wrote {TARGET}")


if __name__ == "__main__":
    main()
